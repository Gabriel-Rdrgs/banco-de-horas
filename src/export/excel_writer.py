"""Exportação para a planilha Excel BancoHoras_Provida_v2.xlsm.

REGRAS CRÍTICAS:
  - BatidasOCR linha 1 = título, linha 2 = cabeçalho, linha 3+ = dados
  - Colunas J(10), K(11), L(12) são FÓRMULAS EXCEL — nunca sobrescrever
  - Sempre abrir com keep_vba=True para preservar macros VBA
  - Salvar com extensão .xlsm (não .xlsx)
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import openpyxl

from src.models.folha import BatidaDiaria, FolhaImportacao
from src.models.log import LogProcessamento
from src.parser.normalizer import time_to_str

logger = logging.getLogger(__name__)

# --- Mapeamento de colunas BatidasOCR (índices 1-based) ---
# ATENÇÃO: J=10, K=11, L=12 são fórmulas — ausentes intencionalmente
_COL_BATIDAS = {
    "ID_BATIDA":            1,   # A
    "COLABORADOR":          2,   # B
    "COMPETENCIA":          3,   # C
    "DATA":                 4,   # D
    "DIA_SEMANA":           5,   # E
    "ENTRADA_MANHA":        6,   # F
    "SAIDA_ALMOCO":         7,   # G
    "VOLTA_ALMOCO":         8,   # H
    "SAIDA_TARDE":          9,   # I
    # J=10 HRS_TRABALHADAS  → FÓRMULA
    # K=11 JORNADA_ESPERADA → FÓRMULA
    # L=12 DIFERENÇA_DIA    → FÓRMULA
    "TIPO_DIA":             13,  # M — preenchido opcionalmente
    "STATUS_OCR":           14,  # N
    "CONFIANCA":            15,  # O
    "OBSERVACOES":          16,  # P
    "DATA_IMPORTACAO":      17,  # Q
    "USUARIO_IMPORTACAO":   18,  # R
}

# --- Mapeamento de colunas ControleOCR (índices 1-based) ---
_COL_CONTROLE = {
    "ID_IMPORTACAO":        1,
    "DATA_HORA":            2,
    "NOME_ARQUIVO":         3,
    "COLABORADOR":          4,
    "COMPETENCIA":          5,
    "LINHAS_IMPORTADAS":    6,
    "CONFIANCA_MEDIA":      7,
    "STATUS":               8,
    "MENSAGEM_ERRO":        9,
    "USUARIO":              10,
}

_HEADER_ROW = 2
_DATA_START = 3


def _proxima_linha_vazia(ws, col: int = 2) -> int:
    """Encontra a primeira linha vazia na coluna especificada, a partir de DATA_START."""
    row = _DATA_START
    while ws.cell(row=row, column=col).value is not None:
        row += 1
    return row


def _id_batida(batida: BatidaDiaria) -> str:
    """Gera ID legível para a batida: COLABORADOR_AAAAMMDD."""
    nome_limpo = batida.colaborador.replace(" ", "_").upper()[:20]
    return f"{nome_limpo}_{batida.data.strftime('%Y%m%d')}"


def _gravar_batidas(ws, batidas: list[BatidaDiaria], usuario: str) -> int:
    """Grava lista de batidas na aba BatidasOCR.

    Não toca em colunas de fórmula (J, K, L).
    Retorna número de linhas gravadas.
    """
    inicio = _proxima_linha_vazia(ws, col=_COL_BATIDAS["COLABORADOR"])
    agora = datetime.now()

    for i, b in enumerate(batidas):
        r = inicio + i
        c = _COL_BATIDAS

        ws.cell(r, c["ID_BATIDA"]).value          = _id_batida(b)
        ws.cell(r, c["COLABORADOR"]).value         = b.colaborador
        ws.cell(r, c["COMPETENCIA"]).value         = b.competencia
        ws.cell(r, c["DATA"]).value                = b.data
        ws.cell(r, c["DIA_SEMANA"]).value          = b.dia_semana
        ws.cell(r, c["ENTRADA_MANHA"]).value       = time_to_str(b.entrada_manha)
        ws.cell(r, c["SAIDA_ALMOCO"]).value        = time_to_str(b.saida_almoco)
        ws.cell(r, c["VOLTA_ALMOCO"]).value        = time_to_str(b.volta_almoco)
        ws.cell(r, c["SAIDA_TARDE"]).value         = time_to_str(b.saida_tarde)
        ws.cell(r, c["STATUS_OCR"]).value          = b.status_ocr
        ws.cell(r, c["CONFIANCA"]).value           = round(b.confianca, 1)
        ws.cell(r, c["OBSERVACOES"]).value         = b.observacoes
        ws.cell(r, c["DATA_IMPORTACAO"]).value     = agora
        ws.cell(r, c["USUARIO_IMPORTACAO"]).value  = usuario

    logger.info(f"BatidasOCR: {len(batidas)} linha(s) gravada(s) a partir da linha {inicio}")
    return len(batidas)


def _gravar_log(ws, log: LogProcessamento) -> None:
    """Grava uma linha de log na aba ControleOCR."""
    r = _proxima_linha_vazia(ws, col=_COL_CONTROLE["ID_IMPORTACAO"])
    c = _COL_CONTROLE

    ws.cell(r, c["ID_IMPORTACAO"]).value     = log.id_importacao
    ws.cell(r, c["DATA_HORA"]).value         = log.data_hora
    ws.cell(r, c["NOME_ARQUIVO"]).value      = log.arquivo
    ws.cell(r, c["COLABORADOR"]).value       = log.colaborador
    ws.cell(r, c["COMPETENCIA"]).value       = log.competencia
    ws.cell(r, c["LINHAS_IMPORTADAS"]).value = log.linhas_importadas
    ws.cell(r, c["CONFIANCA_MEDIA"]).value   = round(log.confianca_media, 1)
    ws.cell(r, c["STATUS"]).value            = log.status
    ws.cell(r, c["MENSAGEM_ERRO"]).value     = log.mensagem_erro
    ws.cell(r, c["USUARIO"]).value           = log.usuario

    logger.info(f"ControleOCR: log {log.id_importacao} gravado na linha {r}")


def exportar_para_excel(
    planilha_path: Path,
    folha: FolhaImportacao,
    log: LogProcessamento,
    dry_run: bool = False,
) -> int:
    """Abre a planilha, grava batidas e log, salva.

    dry_run=True: executa tudo mas não salva — útil para validação.
    Retorna número de batidas gravadas.
    """
    if not planilha_path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {planilha_path}")

    logger.info(f"Abrindo planilha: {planilha_path.name} (dry_run={dry_run})")
    wb = openpyxl.load_workbook(str(planilha_path), keep_vba=True)

    ws_batidas  = wb["BatidasOCR"]
    ws_controle = wb["ControleOCR"]

    n = _gravar_batidas(ws_batidas, folha.batidas, usuario=log.usuario)
    log.linhas_importadas = n
    _gravar_log(ws_controle, log)

    if dry_run:
        logger.warning("dry_run=True — planilha NAO foi salva")
        wb.close()
    else:
        wb.save(str(planilha_path))
        wb.close()
        logger.info(f"Planilha salva: {planilha_path.name}")

    return n
