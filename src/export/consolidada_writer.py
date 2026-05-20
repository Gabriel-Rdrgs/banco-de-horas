"""Grava resultados de apuração na aba Apuracao_Consolidada.

Regras críticas:
  - COLABORADOR deve existir no Cadastro (validação obrigatória)
  - COMPETÊNCIA gravada como 'AAAA-MM' (padrão da planilha)
  - SALDO_ANT = SALDO_LÍQUIDO da competência imediatamente anterior do colaborador
  - SALDO_MÊS = HRS_EXTRAS - FALTAS_ATRASOS  (AJUSTE_MANUAL ignorado)
  - SALDO_LÍQUIDO = SALDO_ANT + SALDO_MÊS
  - Nunca sobrescreve linha existente — cria nova ou avisa se já existe
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl

from src.apuracao import (
    ResultadoMes,
    competencia_para_aaaaMM,
    hhmm_to_td,
    td_to_hhmm,
    nomes_cadastro,
)

logger = logging.getLogger(__name__)

# Mapeamento de colunas da Apuracao_Consolidada (1-based)
# Linha 2 = cabeçalho, linha 3+ = dados
_COL = {
    "NUM":            1,   # A  — #
    "COLABORADOR":    2,   # B
    "SETOR":          3,   # C
    "COMPETENCIA":    4,   # D  — formato AAAA-MM
    "PERIODO":        5,   # E
    "CARGA":          6,   # F  — h/dia
    "TRAB_SAB":       7,   # G
    "SALDO_ANT":      8,   # H
    "HRS_EXTRAS":     9,   # I
    "FALTAS":         10,  # J
    # K (11) = AJUSTE_MANUAL — não gravar
    "SALDO_MES":      12,  # L
    # M (13), N (14) = SALDO_ACUM — calculados por fórmula
    "SALDO_LIQUIDO":  15,  # O
    "FONTE":          16,  # P
    "CRITERIO":       17,  # Q
    "OBSERVACOES":    18,  # R
    "CHAVE":          19,  # S
    "ORIGEM":         24,  # X
    "CONFIABILIDADE": 25,  # Y
}

_HEADER_ROW = 2
_DATA_START = 3


def _primeira_linha_vazia(ws) -> int:
    """Encontra a primeira linha vazia na coluna COLABORADOR (B)."""
    row = _DATA_START
    while ws.cell(row=row, column=_COL["COLABORADOR"]).value is not None:
        row += 1
    return row


def _buscar_linha_existente(ws, colaborador: str, competencia_fmt: str) -> Optional[int]:
    """Retorna índice da linha se já existe entrada para colaborador/competência."""
    row = _DATA_START
    while True:
        collab = ws.cell(row=row, column=_COL["COLABORADOR"]).value
        if collab is None:
            return None
        comp = ws.cell(row=row, column=_COL["COMPETENCIA"]).value
        if str(collab).strip() == colaborador and str(comp or "").strip() == competencia_fmt:
            return row
        row += 1


def _ler_saldo_anterior(ws, colaborador: str, competencia_fmt: str) -> str:
    """Busca SALDO_LÍQUIDO da competência mais recente anterior do colaborador.

    Compara strings 'AAAA-MM' — ordena lexicograficamente (funciona para esse formato).
    Retorna '00:00' se não houver histórico.
    """
    melhor_comp = ""
    melhor_saldo = "00:00"

    row = _DATA_START
    while True:
        collab = ws.cell(row=row, column=_COL["COLABORADOR"]).value
        if collab is None:
            break
        comp_val = str(ws.cell(row=row, column=_COL["COMPETENCIA"]).value or "")
        saldo_val = str(ws.cell(row=row, column=_COL["SALDO_LIQUIDO"]).value or "00:00")

        if (str(collab).strip() == colaborador
                and comp_val < competencia_fmt        # anterior à competência atual
                and comp_val > melhor_comp):           # a mais recente das anteriores
            melhor_comp = comp_val
            melhor_saldo = saldo_val
        row += 1

    if melhor_comp:
        logger.debug(f"Saldo anterior de {colaborador}: {melhor_saldo} (comp {melhor_comp})")
    else:
        logger.debug(f"Sem histórico anterior para {colaborador} — saldo inicial 00:00")

    return melhor_saldo


def gravar_apuracao(
    planilha_path: Path,
    resultado: ResultadoMes,
    dry_run: bool = False,
) -> dict:
    """Grava um ResultadoMes na aba Apuracao_Consolidada.

    Retorna dict com os valores calculados (para exibição no Streamlit).
    Levanta ValueError se colaborador não estiver no Cadastro.
    """
    # Validação: colaborador deve existir no Cadastro
    nomes_validos = nomes_cadastro(planilha_path)
    if resultado.colaborador not in nomes_validos:
        raise ValueError(
            f"Colaborador '{resultado.colaborador}' não encontrado no Cadastro.\n"
            f"Nomes válidos: {nomes_validos}"
        )

    competencia_fmt = competencia_para_aaaaMM(resultado.competencia)   # "2026-04"
    chave = f"{resultado.colaborador}_{competencia_fmt}"

    wb = openpyxl.load_workbook(str(planilha_path), keep_vba=True)
    ws = wb["Apuracao_Consolidada"]

    # Verifica se já existe entrada para este colaborador/competência
    linha_existente = _buscar_linha_existente(ws, resultado.colaborador, competencia_fmt)
    if linha_existente:
        logger.warning(
            f"Já existe entrada para {resultado.colaborador} / {competencia_fmt} "
            f"na linha {linha_existente}. Operação cancelada para evitar duplicata."
        )
        wb.close()
        raise ValueError(
            f"Já existe uma apuração gravada para {resultado.colaborador} "
            f"competência {resultado.competencia}. "
            "Exclua a linha existente antes de reimportar."
        )

    # Busca saldo anterior
    saldo_ant_str = _ler_saldo_anterior(ws, resultado.colaborador, competencia_fmt)
    saldo_ant_td  = hhmm_to_td(saldo_ant_str)
    saldo_liq_td  = saldo_ant_td + resultado.saldo_mes

    # Calcula valores formatados
    vals = {
        "colaborador":    resultado.colaborador,
        "competencia":    competencia_fmt,
        "competencia_br": resultado.competencia,
        "carga":          resultado.carga_horas_dia,
        "trab_sab":       "Sim" if resultado.trabalha_sab else "Não",
        "saldo_ant":      saldo_ant_str,
        "hrs_extras":     td_to_hhmm(resultado.horas_extras),
        "faltas":         td_to_hhmm(resultado.faltas_atrasos),
        "saldo_mes":      td_to_hhmm(resultado.saldo_mes),
        "saldo_liquido":  td_to_hhmm(saldo_liq_td),
        "confiabilidade": round(resultado.confianca_media, 1),
        "chave":          chave,
        "dias_com":       resultado.dias_com_batida,
        "dias_sem":       resultado.dias_sem_batida,
        "dias_uteis":     resultado.total_dias_uteis,
    }

    if dry_run:
        logger.warning(f"dry_run=True — Apuracao_Consolidada NÃO foi alterada")
        wb.close()
        return vals

    # Grava na próxima linha vazia
    r = _primeira_linha_vazia(ws)
    num_linha = r - _DATA_START + 1   # número sequencial da apuração

    c = _COL
    ws.cell(r, c["NUM"]).value          = num_linha
    ws.cell(r, c["COLABORADOR"]).value  = resultado.colaborador
    ws.cell(r, c["COMPETENCIA"]).value  = competencia_fmt
    ws.cell(r, c["CARGA"]).value        = resultado.carga_horas_dia
    ws.cell(r, c["TRAB_SAB"]).value     = vals["trab_sab"]
    ws.cell(r, c["SALDO_ANT"]).value    = saldo_ant_str
    ws.cell(r, c["HRS_EXTRAS"]).value   = vals["hrs_extras"]
    ws.cell(r, c["FALTAS"]).value       = vals["faltas"]
    ws.cell(r, c["SALDO_MES"]).value    = vals["saldo_mes"]
    ws.cell(r, c["SALDO_LIQUIDO"]).value= vals["saldo_liquido"]
    ws.cell(r, c["FONTE"]).value        = "OCR_PYTHON"
    ws.cell(r, c["CRITERIO"]).value     = "✅ Importado via sistema OCR"
    ws.cell(r, c["OBSERVACOES"]).value  = (
        f"Confiança OCR: {resultado.confianca_media:.1f}% | "
        f"Dias c/ batida: {resultado.dias_com_batida} | "
        f"Dias s/ batida: {resultado.dias_sem_batida}"
    )
    ws.cell(r, c["CHAVE"]).value        = chave
    ws.cell(r, c["ORIGEM"]).value       = "Azure Vision OCR"
    ws.cell(r, c["CONFIABILIDADE"]).value = resultado.confianca_media

    wb.save(str(planilha_path))
    wb.close()

    logger.info(
        f"Apuracao_Consolidada: linha {r} gravada — "
        f"{resultado.colaborador} {resultado.competencia} | "
        f"HE={vals['hrs_extras']} FA={vals['faltas']} "
        f"SALDO={vals['saldo_mes']} LIQ={vals['saldo_liquido']}"
    )

    return vals
